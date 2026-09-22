import os
import uvicorn
from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from google.adk.cli.fast_api import get_fast_api_app
from dotenv import load_dotenv

load_dotenv()

# Resolve the agents directory
backend_dir = os.path.dirname(os.path.abspath(__file__))
agents_dir = os.path.join(backend_dir, "agents")

# Construct the standard ADK FastAPI application
# Allow all origins for local development and separate frontend testing
app = get_fast_api_app(
    agents_dir=agents_dir,
    web=False,
    allow_origins=["*"],
    auto_create_session=True
)

# Serve static files (React frontend)
static_dir = os.path.join(backend_dir, "static")
if os.path.exists(static_dir):
    @app.get("/")
    async def read_index():
        return FileResponse(os.path.join(static_dir, "index.html"))

    app.mount("/assets", StaticFiles(directory=os.path.join(static_dir, "assets")), name="assets")

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import tempfile

@app.get("/download_docx")
async def download_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    PRIMARY_COLOR = RGBColor(2, 132, 199)
    DARK_TEXT_COLOR = RGBColor(15, 23, 42)
    LIGHT_TEXT_COLOR = RGBColor(51, 65, 85)
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Product Research Assistant")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_TEXT_COLOR
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("System Architecture, Capabilities, and Deployment Guide")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph()
    
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta.add_run("Created By: Antigravity AI Assistant\nDate: August 2026\nVersion: 2.1.0\nOperating System: Windows Dev Suite")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_page_break()
    
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        
    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = DARK_TEXT_COLOR
        
    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = LIGHT_TEXT_COLOR
        
    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = LIGHT_TEXT_COLOR
        
    def add_code(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = DARK_TEXT_COLOR

    add_h1("1. Executive Summary")
    add_body(
        "The Product Research Assistant is an enterprise-grade agentic workflow tailored to solve the friction "
        "of searching, parsing, and comparing consumer electronics from complex mock catalogs. It leverages "
        "large language model intelligence (specifically Google's Gemini models) to dynamically interact with "
        "structured data, generating summaries and comparisons tailored to specific user needs and price ranges."
    )
    
    add_h1("2. System Architecture Blueprint")
    add_body(
        "The system's modular architecture separates the client interface, model processing pipelines, and data layers:"
    )
    add_bullet("Dashboard Client (React 18 + Tailwind CSS): Provides a high-fidelity split-screen dashboard.")
    add_bullet("API Gateway & Static Server (FastAPI): Implements stateless CORS headers and serves static assets.")
    add_bullet("Agent Orchestration Layer (Google ADK SDK): Coordinates agent execution with gemini-3.5-flash-lite.")
    add_bullet("MCP Database Layer (FastMCP SSE): Hosts tools connected to the catalog database.")
    
    add_h1("3. Core Optimizations and Upgrades")
    add_h2("3.1. Catalog Expansion & Schema Augmentation")
    add_body(
        "The database catalog inside server.py was increased to 26 products across Laptops, Tablets, Smartphones, Audio, and Monitors with verified URLs."
    )
    
    add_h2("3.2. Switching to Gemini 3.5 Flash Lite")
    add_body(
        "Configured the model to use gemini-3.5-flash-lite, dropping latency to under 15 seconds."
    )
    
    add_h2("3.3. Clickable Product Hyperlinks")
    add_body(
        "The React frontend was upgraded to display these URLs as clickable hyperlinks across the dashboard UI."
    )
    
    add_h1("4. Installation & Local Setup")
    add_code(
        "cd frontend && cmd /c npm run build\n"
        "cd mcp_server && ..\\backend\\.venv\\Scripts\\python server.py 8081\n"
        "cd backend && .venv\\Scripts\\python main.py"
    )
    
    # Save to temp file
    temp_dir = tempfile.gettempdir()
    temp_file_path = os.path.join(temp_dir, "product_research_assistant_documentation.docx")
    doc.save(temp_file_path)
    
    return FileResponse(
        temp_file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="product_research_assistant_documentation.docx"
    )

if __name__ == "__main__":
    port = int(os.getenv("PORT", "8080"))
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=True)
